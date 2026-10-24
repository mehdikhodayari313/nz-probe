"""Extract Gachsaran-4 PEP P.PI piping items and match to bank records."""
import openpyxl, json, re, sys, collections
F = '/home/claude/nz-probe/refs/project_files/PEP-Rev.140205_عزيزيان_Gachsaran 4.xlsx'
R = json.load(open('../build/out/records.json', encoding='utf-8'))
for i, r in enumerate(R): r['_i'] = i
V = list(openpyxl.load_workbook(F, read_only=True, data_only=True)['P.PI'].iter_rows(values_only=True))
FM = list(openpyxl.load_workbook(F, read_only=True)['P.PI'].iter_rows(values_only=True))
CLS = {'A': '150', 'C': '300', 'F': '600'}
WT = {  # ASME B36.10 wall thickness mm
 '1/2': dict(STD=2.77, S40=2.77, S80=3.73, S160=4.78, XXS=7.47), '3/4': dict(STD=2.87, S40=2.87, S80=3.91, S160=5.56, XXS=7.82),
 '1': dict(STD=3.38, S40=3.38, S80=4.55, S160=6.35, XXS=9.09), '1 1/2': dict(STD=3.68, S40=3.68, S80=5.08, S160=7.14, XXS=10.15),
 '2': dict(STD=3.91, S40=3.91, S80=5.54, S160=8.74, XXS=11.07), '3': dict(STD=5.49, S40=5.49, S80=7.62, S160=11.13, XXS=15.24),
 '4': dict(STD=6.02, S40=6.02, S80=8.56, S160=13.49, XXS=17.12), '6': dict(STD=7.11, S40=7.11, S80=10.97, S160=18.26, XXS=21.95),
 '8': dict(STD=8.18, S40=8.18, S80=12.70, S160=23.01, S30=7.04, S60=10.31), '10': dict(STD=9.27, S40=9.27, S80=15.09, S160=28.58, S30=7.80, S60=12.70),
 '12': dict(STD=9.53, S40=10.31, S80=17.48, S160=33.32, S30=8.38, S60=14.27), '14': dict(STD=9.53, S40=11.13, S80=19.05, S30=9.53, S60=15.09),
 '16': dict(STD=9.53, S40=12.70, S80=21.44, S30=9.53, S60=16.66), '18': dict(STD=9.53, S40=14.27, S80=23.83, S30=11.13),
 '20': dict(STD=9.53, S40=15.09, S80=26.19, S30=12.70), '24': dict(STD=9.53, S40=17.48, S80=30.96, S30=14.27)}


def sz(v):
    s = str(v).replace('"', '').strip()
    m = re.fullmatch(r'(\d)(\d)/(\d)', s)          # 11/2 -> 1 1/2
    if m: s = f'{m.group(1)} {m.group(2)}/{m.group(3)}'
    if s.endswith('.0'): s = s[:-2]
    return s


def first(s): return str(s).split(' / ')[0].split(' x ')[0].strip()


def find(fam, size, rating, item=None, pred=None, prefer=None):
    c = [r for r in R if r['fam'] == fam and first(r['size']) == size and r['rating'] == rating and r['price'] is not None
         and (item is None or r['item'].startswith(item)) and (pred is None or pred(r))]
    if prefer and len(c) > 1:
        p = [r for r in c if prefer(r)]; c = p or c
    return c[0] if c else None


def sch(s):
    s = str(s or '').upper().replace('SCH.SCH.', 'SCH.').replace('SCH.', '').strip()
    return s


out = []
for v, fm in zip(V[3:], FM[3:]):
    if not v[0] or not isinstance(v[12], (int, float)) or v[6] == 'Lot':
        continue
    m = re.match(r'=([\d.]+)\*USD_Inflation_Factor', str(fm[12]))
    if not m: continue
    usd92 = float(m.group(1)); usd02 = v[12]
    item, typ, b1, b2, unit, cls, sc, mat, rem = v[1], v[3], sz(v[4]), sz(v[5]), v[6], v[8], sch(v[9]), str(v[10] or '').strip(), v[11]
    rating = CLS.get(str(cls)[0]) if cls else None
    nonmetal = mat in ('GRE', 'HDPE 100', 'ASTM B148') or str(cls).startswith(('AX', 'AZ', 'FX'))
    rec = None; note = ''
    if nonmetal:
        note = 'غیرفلزی/برنز یا خط غیرفلزی — مقایسه نشد'
    elif item == 'Valve' and rating:
        fam = {'Gate': 'GAT', 'Globe': 'GLB', 'Glob': 'GLB', 'Check': 'CHK', 'Ball': 'BAL'}.get(typ)
        forged = 'A105' in mat
        rec = find(fam, b1, f'{rating} RF', prefer=(lambda r: r['remark'].startswith('F')) if forged else (lambda r: r['remark'].startswith('C')))
    elif item == 'Flange' and typ == 'Flange' and rating:
        rec = find('FLG', b1, rating, item='FLANGE W.N'); note = 'نوع فلنج در PEP ذکر نشده؛ WN فرض شد'
    elif item == 'Flange' and typ == 'Spectace' and rating:
        rec = find('SPB', b1, rating)
    elif item == 'Special' and typ == 'Orifice Flange' and rating:
        rec = find('FLG', b1, rating, item='FLANGE ORIFICE')
    elif item == 'Fitting':
        bs = {'STD': 'SCH 40', '40': 'SCH 40', '80': 'SCH 80', '160': 'SCH 160', '30': 'SCH 30', '20': 'SCH 20'}.get(sc)
        if typ in ('90deg.Elbow', '45deg.Elbow', 'Cap', 'Tee', 'Reducer') and bs:
            fam = {'90deg.Elbow': 'E90', '45deg.Elbow': 'E45', 'Cap': 'CAP', 'Reducer': 'CRD'}.get(typ)
            if typ == 'Tee':
                fam = 'TEE' if b2 in ('-', b1) else 'TRD'
            rec = find(fam, b1, bs)
            if fam in ('TRD', 'CRD') and rec and rec['size'] != f'{b1} x {b2}':
                note = f'کاهش سایز بانک {rec["size"]} ≠ {b1} x {b2}'
            if sc == 'STD': note = (note + ' | ' if note else '') + 'STD ≈ SCH 40'
        elif typ == 'Weldolet':
            r_ = {'STD': 'STANDARD WEIGHT (REDUCING SIZE (STD))', '40': 'STANDARD WEIGHT (REDUCING SIZE (STD))', '80': 'EXTRA STRONG RED. SIZE (XS)',
                  '160': 'SCH 160 & DOUBLE EXTRA STRONG RED. SIZE (SCH160 & XXS)', 'XXS': 'SCH 160 & DOUBLE EXTRA STRONG RED. SIZE (SCH160 & XXS)'}.get(sc)
            if r_: rec = find('WOL', b2, r_)
        elif typ == 'Sockolet':
            rec = find('OLT', b2, '3000', item='SOCKOLET'); note = 'کلاس ۳۰۰۰ فرض شد'
    elif item == 'Gasket' and mat == 'SPIRAL WOUND' and rating:
        rec = find('SPW', b1, rating)
    elif item == 'StudBolt&nuts' and isinstance(v[5], (int, float)):
        rec = find('BLT', b1, f'L = {int(v[5])} mm', item='STUD BOLT')
    elif item == 'Pipe':
        k = 'STD' if sc == 'STD' else 'XXS' if sc == 'XXS' else 'S' + sc
        wt = WT.get(b1, {}).get(k)
        if wt:
            cands = [r for r in R if r['fam'] == 'PIPE' and first(r['size']) == b1 and r['price'] is not None]
            best = min(cands, key=lambda r: abs(float(re.search(r'WT ([\d.]+)', r['rating']).group(1)) - wt), default=None)
            if best and abs(float(re.search(r'WT ([\d.]+)', best['rating']).group(1)) - wt) <= 0.25:
                rec = best; note = f'WT {wt} mm از ASME B36.10؛ بانک ERW، قیمت بانک به ازای متر'
    bank_k = None
    if rec:
        bank_k = rec['price']
        if item == 'Pipe':
            L = float(re.search(r'L ([\d.]+)', rec['spec']).group(1)); bank_k = rec['price'] / L
    out.append(dict(no=v[0], item=item, typ=typ, b1=b1, b2=b2, cls=cls, rating=rating, sch=sc, mat=mat, nace=rem, qty=v[7], unit=unit,
                    usd92=usd92, usd02=usd02, fam=rec['fam'] if rec else '', rid=rec['_i'] if rec else None, bank_k=bank_k, note=note))
json.dump(out, open('g4.json', 'w', encoding='utf-8'), ensure_ascii=False, indent=0)
print('rows', len(out), 'matched', sum(1 for o in out if o['rid'] is not None))
print(collections.Counter((o['item'], o['typ'], o['fam']) for o in out if o['rid'] is not None))
print(collections.Counter((o['item'], o['typ']) for o in out if o['rid'] is None and not o['note'].startswith('غیر')))
