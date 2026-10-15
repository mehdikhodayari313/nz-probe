# اجرا: python3 -I tool/make_compressor_fixture.py <Air_Compressor_Price_Bank_v8.x.xlsx> test/fixtures_private/Compressors_bank_test.xlsx
"""نسخه‌ی کوچک بانک کمپرسور هوا برای آزمون: سلول‌های ورودی Master، فروشندگان، استعلام‌ها، Test_Cases و
مقادیر ستون‌های فرمولی Master (شیت Expected؛ همه مقدار، بدون فرمول)."""
import sys, openpyxl
from openpyxl.styles import PatternFill
from openpyxl.comments import Comment

src, dst = sys.argv[1], sys.argv[2]
wv = openpyxl.load_workbook(src, data_only=True)
wf = openpyxl.load_workbook(src)
out = openpyxl.Workbook(); out.remove(out.active)
m = out.create_sheet('Master Database'); sv = wv['Master Database']; sf = wf['Master Database']
e = out.create_sheet('Expected')
m['A1'] = sv['A1'].value
ji = je = 0
for c in range(1, sv.max_column + 1):
    h = sv.cell(4, c).value
    formula = isinstance(sf.cell(5, c).value, str) and sf.cell(5, c).value.startswith('=')
    if formula:
        je += 1; e.cell(1, je, h)
        for r in range(5, sv.max_row + 1):
            if sv.cell(r, 1).value is not None: e.cell(r - 3, je, sv.cell(r, c).value)
        continue
    ji += 1; m.cell(4, ji, h)
    for r in range(5, sv.max_row + 1):
        if sv.cell(r, 1).value is None: continue
        x = m.cell(r, ji, sv.cell(r, c).value)
        rgb = sf.cell(r, c).fill.fgColor.rgb if sf.cell(r, c).fill else None
        if rgb in ('FFFFFF00', 'FFDDEBF7'): x.fill = PatternFill('solid', fgColor=rgb)
        if sf.cell(r, c).comment: x.comment = Comment(sf.cell(r, c).comment.text, 'bank')
for name in ['Vendor_Register', 'Vendor_Quotes', 'Test_Cases']:
    s, d = wv[name], out.create_sheet(name)
    for row in s.iter_rows(min_row=1, max_row=s.max_row):
        for c in row:
            if c.value is not None: d.cell(c.row, c.column, c.value)
out.save(dst)
