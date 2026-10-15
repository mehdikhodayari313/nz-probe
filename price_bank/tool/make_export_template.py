"""ساخت شیت «خروجی_قطعی» نمونه — برای کپی در فایل‌های اکسل تخمین.

اجرا:  python3 tool/make_export_template.py <out.xlsx>
"""
import sys

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "خروجی_قطعی"
ws.sheet_view.rightToLeft = True

head = ["کد یکتا", "نوع", "رده", "عنوان", "مقدار / فرمول / محدوده", "واحد",
        "وضعیت", "نسخه", "توضیح / منبع", "ورودی‌ها (فقط آزمون)"]
widths = [22, 10, 12, 34, 30, 10, 14, 8, 40, 34]
for i, (h, w) in enumerate(zip(head, widths), 1):
    c = ws.cell(1, i, h)
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor="1F4E78")
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.column_dimensions[c.column_letter].width = w
ws.freeze_panes = "A2"

samples = [
    ["PUMP.MAT.S-5", "ضریب", "پمپ", "ضریب کلاس جنس S-5", 1, "-", "قطعی", 1, "مبنای ضرایب جنس"],
    ["PUMP.EXP.N", "ضریب", "پمپ", "توان n مقیاس‌دهی با توان", 0.62, "-", "در حال بررسی", 1, "از برازش ۱۴۶ رکورد"],
    ["VES.WEIGHT.K", "فرمول", "ظرف", "ضریب وزن ظرف", "=1*1", "-", "در حال بررسی", 1,
     "به‌جای ‎=1*1‎ فرمول واقعی را بنویسید؛ سلول‌های ارجاعی خودکار خوانده می‌شوند"],
    ["GEN.CLASS.FACTORS", "جدول", "عمومی", "جدول ضرایب کلاس فشار", "'ضرایب'!A1:D12", "-", "در حال بررسی", 1,
     "آدرس محدوده را به صورت متن بنویسید"],
    ["PUMP.CASE.01", "آزمون", "پمپ", "پمپ سانتریفیوژ نمونه — قیمت تعدیل‌شده ۱۴۰۴", 123456, "EUR", "در حال بررسی", 1,
     "عدد مورد انتظار از همین اکسل", "flow=120; head=80; power=55; year=1404"],
]
for r, row in enumerate(samples, 2):
    for c, v in enumerate(row, 1):
        ws.cell(r, c, v)

kind = DataValidation(type="list", formula1='"ضریب,فرمول,جدول,شاخص,آزمون,قاعده"', allow_blank=True)
status = DataValidation(type="list", formula1='"قطعی,در حال بررسی"', allow_blank=True)
ws.add_data_validation(kind)
ws.add_data_validation(status)
kind.add("B2:B2000")
status.add("G2:G2000")

g = wb.create_sheet("راهنما")
g.sheet_view.rightToLeft = True
g.column_dimensions["A"].width = 110
for i, t in enumerate([
    "راهنمای شیت «خروجی_قطعی»",
    "۱. این شیت را در هر فایل اکسل تخمین کپی کنید (نام شیت دقیقاً «خروجی_قطعی» بماند).",
    "۲. هر نتیجه‌ای که برنامه باید داشته باشد یک ردیف است. فقط ردیف‌های با وضعیت «قطعی» به برنامه می‌رسند.",
    "۳. «کد یکتا» را بعداً عوض نکنید؛ تغییرات نسخه‌به‌نسخه با همین کد شناخته می‌شوند. اگر عدد عوض شد، «نسخه» را یکی بالا ببرید.",
    "۴. ضریب/شاخص: عدد. فرمول: خود فرمول اکسل (سلول‌های ارجاعی و مقدارشان خودکار خوانده می‌شوند).",
    "۵. جدول: آدرس محدوده به صورت متن، مثل ‎'ضرایب'!A1:D12‎ — سطر اول محدوده عنوان ستون‌ها باشد.",
    "۶. آزمون: ورودی‌ها در ستون J (مثل ‎flow=120; head=80‎) و عدد درست در ستون E. برنامه باید دقیقاً همین عدد را بدهد.",
    "۷. قاعده: یک قاعده‌ی متنی (مثلاً «اگر NACE باشد ضریب ۱٫۱۵ اعمال شود») در ستون E یا I.",
    "۸. پیش از ارسال، فایل را یک بار در اکسل ذخیره کنید تا مقدار فرمول‌ها در فایل باشد.",
], 1):
    c = g.cell(i, 1, t)
    if i == 1:
        c.font = Font(bold=True, size=13)

wb.save(sys.argv[1])
print("ok", sys.argv[1])
