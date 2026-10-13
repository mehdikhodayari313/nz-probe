"""تولید داده‌های پایه‌ی برنامه (assets/seed/*.json) از فایل‌های مرجع.

اجرا:  python3 tool/gen_seed.py <Pumps_Price_Bank_v8.xlsx> <Pump_Escalation_Validation_2.xlsx> <out_dir>
خروجی شامل اطلاعات قیمت تجهیزات نیست؛ فقط ساختار، فهرست‌ها، ضرایب و شاخص‌های عمومی.
"""
import json, sys, openpyxl

v8_path, val_path, out = sys.argv[1], sys.argv[2], sys.argv[3]
wb = openpyxl.load_workbook(v8_path)
wv = openpyxl.load_workbook(v8_path, data_only=True)
vb = openpyxl.load_workbook(val_path, data_only=True)

def dump(name, obj):
    with open(f"{out}/{name}", "w", encoding="utf-8") as f:
        json.dump(obj, f, ensure_ascii=False, indent=1)

# ---------------------------------------------------------------- lists
lc = wv["Lookup_Codes"]
lists = {}
for col in range(1, lc.max_column + 1):
    key = lc.cell(4, col).value
    if not key or str(key).startswith("SRCH_"):
        continue
    vals = []
    for r in range(5, lc.max_row + 1):
        v = lc.cell(r, col).value
        if v is None:
            continue
        vals.append(str(v) if not isinstance(v, (int, float)) or isinstance(v, bool) else str(int(v)) if float(v).is_integer() else str(v))
    lists[key] = vals

# فارسی زیرگروه‌ها از داشبورد
dash = wv["Executive Dashboard"]
subcat_fa = {}
for r in range(8, 14):
    k = dash.cell(r, 1).value
    if k:
        subcat_fa[k] = {"name_en": dash.cell(r, 2).value, "desc_fa": dash.cell(r, 3).value}

# جدول‌های مرجع داشبورد (کلاس متریال، نوع API 610، Incoterm)
def table(start_title, ncols):
    rows, on = [], False
    for r in range(1, dash.max_row + 1):
        a = dash.cell(r, 1).value
        if isinstance(a, str) and a.startswith(start_title):
            on = True
            continue
        if on:
            if a is None or (isinstance(a, str) and a[:2].strip() in ("2)", "3)", "4)", "5)")):
                if rows:
                    break
                continue
            rows.append([dash.cell(r, c).value for c in range(1, ncols + 1)])
    return rows[1:]  # بدون سرستون

ref_tables = {
    "material_classes": table("2)", 5),
    "api610_types": table("3)", 3),
    "incoterms": table("4)", 4),
}
dump("lists.json", {"lists": lists, "subcategories": subcat_fa, "reference_tables": ref_tables})

# ---------------------------------------------------------------- pump template
md = wb["Master Database"]
group_of = {}
groups = [("A", "E", "identification", "شناسایی"), ("F", "U", "key_specs", "مشخصات کلیدی"),
          ("V", "X", "extra_specs", "مشخصات تکمیلی"), ("Y", "AI", "price", "قیمت"),
          ("AJ", "AK", "source", "منبع و سازنده"), ("AL", "AR", "scope", "دامنه کار (Scope)"),
          ("AS", "AS", "control", "کنترل و ردیابی"), ("AT", "BD", "calc", "محاسبات نرمال‌سازی و تعدیل"),
          ("BE", "BG", "control", "کنترل و ردیابی")]
from openpyxl.utils import column_index_from_string as ci, get_column_letter as gl
for a, b, k, fa in groups:
    for c in range(ci(a), ci(b) + 1):
        group_of[gl(c)] = (k, fa)

# key, col, label_fa, storage, type, list, unit
F = [
 ("v8_row","A","ردیف v8","core","number",None,None),
 ("name","B","نام تجهیز / سرویس","core","text",None,None),
 ("project","C","پروژه","core","text",None,None),
 ("tag","D","شماره تگ","core","text",None,None),
 ("sub_category","E","زیرگروه","spec","enum","SubCat",None),
 ("api610_type","F","نوع API 610","spec","enum","API610",None),
 ("material_class","G","کلاس متریال API 610","spec","enum","MAT",None),
 ("capacity_m3h","H","دبی","spec","number",None,"m³/h"),
 ("diff_head_m","I","هد اختلافی","spec","number",None,"m"),
 ("power_kw","J","توان","spec","number",None,"kW"),
 ("size_kw","K","توان مبنای اندازه (برآوردی)","calc","number",None,"kW"),
 ("standards","L","استانداردها","calc","text",None,None),
 ("pump_standard","M","استاندارد پمپ","spec","enum","Std",None),
 ("ips_standard","N","استاندارد IPS","spec","enum","IPS",None),
 ("nace","O","NACE MR0175","spec","enum","YNS",None),
 ("seal_arrangement","P","نوع / آرایش آب‌بند","spec","enum","SEAL_ARR",None),
 ("api682_category","Q","گروه API 682","spec","enum","SEAL_CAT",None),
 ("api682_plan","R","پلن آب‌بندی API 682","spec","enum","SEAL_PLAN",None),
 ("driver_type","S","نوع محرک","spec","enum","DRV",None),
 ("service_fluid","T","سیال سرویس","spec","enum","Service",None),
 ("quantity","U","تعداد","core","number",None,"دستگاه"),
 ("pump_family","V","خانواده پمپ","spec","enum","Family",None),
 ("discharge_pressure_barg","W","فشار رانش","spec","number",None,"barg"),
 ("power_basis","X","مبنای توان","spec","enum","Power_Basis",None),
 ("cost_year","Y","سال قیمت (شمسی)","price","enum","Year",None),
 ("currency","Z","ارز مرجع","price","enum","Currency",None),
 ("unit_price","AA","قیمت واحد (ارز مرجع)","price","number",None,None),
 ("eur_source","AB","قیمت یورو (سال منبع)","calc","number",None,"EUR"),
 ("usd_source","AC","قیمت دلار (سال منبع)","calc","number",None,"USD"),
 ("irr_source","AD","قیمت ریالی (سال منبع)","calc","number",None,"IRR"),
 ("price_basis","AE","مبنای قیمت (طبق منبع)","price","text",None,None),
 ("incoterm","AF","شرط تحویل (Incoterm)","price","enum","Incoterm",None),
 ("eur_target","AG","قیمت یورو تعدیل‌شده (سال هدف)","calc","number",None,"EUR"),
 ("usd_target","AH","قیمت دلار تعدیل‌شده (سال هدف)","calc","number",None,"USD"),
 ("irr_target","AI","قیمت ریالی تعدیل‌شده (سال هدف)","calc","number",None,"IRR"),
 ("manufacturer","AJ","سازنده / تأمین‌کننده (طبق منبع)","spec","text",None,None),
 ("project_phase","AK","فاز پروژه","spec","enum","Phase",None),
 ("scope_driver","AL","Scope: محرک","spec","enum","YNS",None),
 ("scope_baseplate","AM","Scope: شاسی و کوپلینگ","spec","enum","YNS",None),
 ("scope_seal_support","AN","Scope: سیستم پشتیبان آب‌بند","spec","enum","YNS",None),
 ("scope_comm_spares","AO","Scope: یدکی راه‌اندازی","spec","enum","YNS",None),
 ("scope_2yr_spares","AP","Scope: یدکی دوساله","spec","enum","YNS",None),
 ("scope_tpi","AQ","Scope: بازرسی شخص ثالث","spec","enum","YNS",None),
 ("scope_fat","AR","Scope: آزمون کارخانه با حضور","spec","enum","YNS",None),
 ("qa_flags","AS","هشدارهای کیفیت داده","calc","text",None,None),
 ("mat_f","AT","ضریب متریال","calc","number",None,None),
 ("seal_f","AU","ضریب آب‌بند","calc","number",None,None),
 ("drive_f","AV","ضریب محرک","calc","number",None,None),
 ("scope_f","AW","ضریب Scope","calc","number",None,None),
 ("all_f","AX","ضریب کل پیکربندی","calc","number",None,None),
 ("normalized_eur","AY","یورو نرمال‌شده (پیکربندی پایه، سال منبع)","calc","number",None,"EUR"),
 ("ppi_base","AZ","PPI سال پایه","calc","number",None,None),
 ("ppi_target","BA","PPI سال هدف","calc","number",None,None),
 ("floor_applies","BB","مشمول کف قیمت","calc","text",None,None),
 ("normalized_escalated_eur","BC","یورو نرمال‌شده‌ی تعدیل‌شده","calc","number",None,"EUR"),
 ("cost_coef","BD","ضریب هزینه (یورو ÷ kWⁿ)","calc","number",None,None),
 ("source_document","BE","سند منبع","price","text",None,None),
 ("remarks","BF","ملاحظات","core","text",None,None),
 ("include_in_totals","BG","شمارش در جمع‌ها","core","enum","YN",None),
]
fields = []
for i, (key, col, fa, storage, typ, lst, unit) in enumerate(F):
    hdr = md[f"{col}4"]
    gk, gfa = group_of[col]
    fields.append({
        "key": key, "v8_column": col, "label_fa": fa, "label_en": hdr.value,
        "group": gk, "group_fa": gfa, "storage": storage, "type": typ,
        "list": lst, "unit": unit, "sort": i,
        "help_fa": (hdr.comment.text.strip() if hdr.comment else None),
    })
dump("pump_template.json", {"category": "PUMP", "fields": fields})

# ---------------------------------------------------------------- factors
fr = wv["Factors_ROT"]
factors = []
for r in range(5, fr.max_row + 1):
    g, o = fr.cell(r, 2).value, fr.cell(r, 3).value
    if not g:
        continue
    factors.append({"group": g, "option": str(o), "provisional": fr.cell(r, 4).value,
                    "calibrated": fr.cell(r, 5).value, "description": fr.cell(r, 8).value})
dump("factors.json", {"set": "PUMP_ROT", "base_config": fr["A1"].value, "factors": factors})

# ---------------------------------------------------------------- indices
series, values = [], []
def S(code, name_fa, name_en, kind, group, unit, base, source, calendar="jalali"):
    series.append(dict(code=code, name_fa=name_fa, name_en=name_en, kind=kind, group=group,
                       unit=unit, base=base, source=source, calendar=calendar))
def V(code, period, value, status="official", source=None):
    if value is None:
        return
    values.append(dict(series=code, period=str(period), value=float(value), status=status, source=source))

ind = vb["Indices"]
S("PPI_EU_C2813_GREG", "PPI پمپ و کمپرسور منطقه‌ی یورو (میلادی)", "Eurostat PPI C28.13 EA20, Gregorian",
  "level", "ppi_world", "2021=100", "2021=100", "Eurostat sts_inppd_a, NACE C28.13", "gregorian")
for r in range(10, 33):
    y, v, st = ind.cell(r, 1).value, ind.cell(r, 2).value, ind.cell(r, 4).value or ""
    V("PPI_EU_C2813_GREG", y, v, "forecast" if "پیش" in st else ("imputed" if "imputed" in st else "official"), st)
S("PPI_US_333914_GREG", "PPI پمپ آمریکا (میلادی)", "US BLS PPI NAICS 333914 (January)",
  "level", "ppi_world", "Dec1983=100", "Dec1983=100", "FRED/BLS PCU333914333914", "gregorian")
for r in range(10, 33):
    V("PPI_US_333914_GREG", ind.cell(r, 1).value, ind.cell(r, 5).value)

S("PPI_EU_C2813", "PPI پمپ منطقه‌ی یورو (سال شمسی)", "Eurostat PPI C28.13 EA20, Jalali-year blend",
  "level", "ppi_world", "2021=100", "2021=100",
  "ترکیب سال میلادی: 0.79×PPI(Y+621) + 0.21×PPI(Y+622) — Pump_Escalation_Validation, Indices C51:C71")
esc = wv["Escalation_Index"]
for r in range(5, 26):
    st = esc.cell(r, 4).value or ""
    V("PPI_EU_C2813", esc.cell(r, 1).value, esc.cell(r, 2).value,
      "forecast" if "FORECAST" in st or "forecast" in st else "official", st)

# نرخ ارز — نوع «جدول v8»
S("FX_USD_V8", "دلار به ریال — جدول مرجع v8", "USD→IRR (v8 reference table)", "level", "fx", "ریال", None,
  "Pumps_Price_Bank_v8, Assumptions & FX Reference")
S("FX_EUR_V8", "یورو به ریال — جدول مرجع v8", "EUR→IRR (v8 reference table)", "level", "fx", "ریال", None,
  "Pumps_Price_Bank_v8, Assumptions & FX Reference")
af = wv["Assumptions & FX Reference"]
for r in range(5, 26):
    y = af.cell(r, 1).value
    note = af.cell(r, 5).value
    V("FX_USD_V8", y, af.cell(r, 2).value, "official", note)
    V("FX_EUR_V8", y, af.cell(r, 3).value, "official", note)

# نرخ ارز — بازار آزاد (فایل صحت‌سنجی)
S("FX_USD_FREE", "دلار به ریال — بازار آزاد", "USD→IRR free market (annual avg, approx.)", "level", "fx", "ریال",
  None, "Pump_Escalation_Validation, Indices ستون D")
S("FX_EUR_FREE", "یورو به ریال — بازار آزاد", "EUR→IRR free market (effective)", "level", "fx", "ریال",
  None, "Pump_Escalation_Validation, Indices ستون G (= دلار آزاد × یورو/دلار، یا نرخ جایگزین)")
S("FX_USD_NIMA", "دلار به ریال — مرکز مبادله / نیما", "USD→IRR exchange centre (NIMA)", "level", "fx", "ریال",
  None, "ورود توسط کاربر")
S("FX_EUR_NIMA", "یورو به ریال — مرکز مبادله / نیما", "EUR→IRR exchange centre (NIMA)", "level", "fx", "ریال",
  None, "ورود توسط کاربر")
S("INF_IRR", "تورم ریالی سالانه (CPI)", "Rial inflation (annual)", "annual_rate", "inflation", "نسبت", None,
  "مارون ۳ (شیت Inflation) / Pump_Escalation_Validation")
S("INF_FX_M3", "تورم ارزی سالانه (مارون ۳)", "FX inflation (Maroon 3)", "annual_rate", "inflation", "نسبت", None,
  "مارون ۳ (شیت Inflation)")
for r in range(51, 72):
    y = ind.cell(r, 1).value
    st = ind.cell(r, 13).value
    V("FX_USD_FREE", y, ind.cell(r, 4).value, "approx", st)
    V("FX_EUR_FREE", y, ind.cell(r, 7).value, "approx",
      ("نرخ جایگزین (مارون ۳)" if ind.cell(r, 6).value else st))
    V("INF_IRR", y, ind.cell(r, 8).value, "official")
    V("INF_FX_M3", y, ind.cell(r, 10).value, "official")

# فلزات
metals = [("MET_STEEL", "فولاد (HRC/Plate)", 38), ("MET_NICKEL", "نیکل", 39),
          ("MET_MOLY", "مولیبدن", 40), ("MET_COPPER", "مس", 41)]
for code, fa, r in metals:
    S(code, f"قیمت {fa}", code, "level", "metal", "USD/t", None,
      f"Pump_Escalation_Validation, Indices ردیف {r}؛ وزن در سبد پمپ: {ind.cell(r, 2).value}")
    for c, y in zip(range(3, 7), [1393, 1403, 1404, 1405]):
        V(code, y, ind.cell(r, c).value, "approx")

# PPI بخشی ایران (مرکز آمار، 1395=100)
iran = [("PPI_IR_TOTAL", "PPI کل بخش صنعت ایران", 1758.0),
        ("PPI_IR_C24", "PPI ساخت فلزات پایه (ISIC 24)", 2242.8),
        ("PPI_IR_C25", "PPI ساخت محصولات فلزی ساخته‌شده (ISIC 25)", 2011.8),
        ("PPI_IR_C26", "PPI محصولات رایانه‌ای، الکترونیکی و نوری (ISIC 26)", 1000.4),
        ("PPI_IR_C27", "PPI ساخت تجهیزات برقی (ISIC 27)", 2715.7),
        ("PPI_IR_C28", "PPI ساخت ماشین‌آلات و تجهیزات (ISIC 28)", 2439.3)]
for code, fa, v in iran:
    S(code, fa, code, "level", "ppi_iran", "1395=100", "1395=100", "مرکز آمار ایران — شاخص قیمت تولیدکننده بخش صنعت")
    V(code, "1404-03", v, "needs_verification",
      "https://amar.org.ir/Portals/0/Articles/sanaat-3-1404.pdf — خوانده‌شده خودکار؛ با PDF تطبیق دهید")

dump("indices.json", {"series": series, "values": values})

# ---------------------------------------------------------------- settings
dump("settings.json", {
    "target_year": 1404, "fx_rate_type": "V8", "ppi_series": "PPI_EU_C2813",
    "floor_eur": 25000, "floor_base_year": 1404, "cost_exponent_n": 0.6, "pump_efficiency": 0.6,
    "eur_check_limit": 5000000, "display_unit": "IRR",
    "fx_rate_types": [{"code": "V8", "name_fa": "جدول مرجع v8"},
                      {"code": "FREE", "name_fa": "بازار آزاد"},
                      {"code": "NIMA", "name_fa": "مرکز مبادله / نیما"}],
})
print("series", len(series), "values", len(values), "fields", len(fields), "factors", len(factors),
      "lists", len(lists))
