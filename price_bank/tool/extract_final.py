"""استخراج «موارد قطعی» از فایل‌های اکسل تخمین به خلاصه‌ی کوچک JSON.

هدف: Claude برای ارتقای برنامه به‌جای خواندن کل اکسل، فقط همین خلاصه را بخواند
(کم‌هزینه و بدون خطای برداشت).

قرارداد اکسل: یک شیت به نام «خروجی_قطعی» (یا EXPORT) با این ستون‌ها از ردیف ۲ به بعد:
  A کد یکتا        مثلاً PUMP.MAT.I-1  (ثابت بماند؛ مبنای مقایسه‌ی نسخه‌هاست)
  B نوع            ضریب | فرمول | جدول | شاخص | آزمون | قاعده
  C رده            پمپ، ظرف، مبدل، ... یا «عمومی»
  D عنوان
  E مقدار          عدد، فرمول اکسل، یا برای «جدول» آدرس محدوده مثل 'ضرایب'!A1:F30
  F واحد
  G وضعیت          فقط «قطعی» استخراج می‌شود
  H نسخه
  I توضیح / منبع
  J ورودی‌ها       (برای «آزمون»): flow=120; head=80; year=1404

اجرا:
  python3 -I tool/extract_final.py <file.xlsx> [--spec spec] [--family نام] [--force]
خروجی:
  spec/pending/<تاریخ>_<نام فایل>.json  فقط موارد جدید/تغییرکرده/حذف‌شده
  spec/pending/<تاریخ>_<نام فایل>.md    خلاصه‌ی فارسی برای شما
  spec/state.json                       آخرین وضعیت استخراج‌شده (مبنای مقایسه‌ی بعدی)
"""
import datetime
import hashlib
import json
import os
import re
import sys

import openpyxl
from openpyxl.utils import range_boundaries

SHEET_NAMES = ("خروجی_قطعی", "خروجی قطعی", "EXPORT", "Export", "export")
FINAL_WORDS = {"قطعی", "final", "Final", "FINAL"}
KINDS = {"ضریب", "فرمول", "جدول", "شاخص", "آزمون", "قاعده"}
MAX_TABLE_CELLS = 5000


def clean(v):
    if isinstance(v, float) and v.is_integer():
        return int(v)
    if isinstance(v, (datetime.date, datetime.datetime)):
        return v.isoformat()
    if isinstance(v, str):
        return v.strip()
    return v


def find_sheet(wb):
    for n in SHEET_NAMES:
        if n in wb.sheetnames:
            return n
    return None


def refs_of(formula, home_sheet):
    """ارجاع‌های یک فرمول: [(sheet, range_text)]"""
    out = []
    body = formula[1:] if formula.startswith("=") else formula
    body = re.sub(r'"[^"]*"', "", body)  # متن‌های داخل گیومه
    for m in re.finditer(r"(?<![\w.])(?:(?:'(?:[^']|'')+'|[^\s\(\),;=+\-*/^&<>!:'\"]+)!)?\$?[A-Z]{1,3}\$?\d+(?::\$?[A-Z]{1,3}\$?\d+)?(?![\w(])", body):
        tok = m.group(0)
        if "!" in tok:
            sh, rng = tok.rsplit("!", 1)
            sh = sh.strip("'").replace("''", "'")
        else:
            sh, rng = home_sheet, tok
        out.append((sh, rng.replace("$", "")))
    return out


def read_range(wv, sheet, rng):
    if sheet not in wv.sheetnames:
        return None, f"شیت «{sheet}» پیدا نشد"
    ws = wv[sheet]
    try:
        c1, r1, c2, r2 = range_boundaries(rng)
    except ValueError:
        return None, f"آدرس «{rng}» نامعتبر است"
    if (r2 - r1 + 1) * (c2 - c1 + 1) > MAX_TABLE_CELLS:
        return None, f"محدوده‌ی «{rng}» بیش از {MAX_TABLE_CELLS} سلول است"
    rows = []
    for r in range(r1, r2 + 1):
        rows.append([clean(ws.cell(r, c).value) for c in range(c1, c2 + 1)])
    while rows and all(v is None for v in rows[-1]):
        rows.pop()
    return rows, None


def parse_inputs(text):
    out = {}
    if not text:
        return out
    for part in re.split(r"[;؛\n]", str(text)):
        if "=" in part:
            k, v = part.split("=", 1)
            v = v.strip()
            try:
                v = clean(float(v.replace(",", "")))
            except ValueError:
                pass
            out[k.strip()] = v
    return out


def extract(path):
    wf = openpyxl.load_workbook(path, data_only=False)
    wv = openpyxl.load_workbook(path, data_only=True)
    name = find_sheet(wf)
    if not name:
        raise SystemExit(f"خطا: شیت «خروجی_قطعی» در {os.path.basename(path)} نیست.")
    sf, sv = wf[name], wv[name]
    items, warnings = {}, []
    for r in range(2, sf.max_row + 1):
        code = clean(sv.cell(r, 1).value)
        if not code:
            continue
        status = clean(sv.cell(r, 7).value)
        if status not in FINAL_WORDS:
            continue
        kind = clean(sv.cell(r, 2).value)
        if kind not in KINDS:
            warnings.append(f"ردیف {r} ({code}): نوع «{kind}» شناخته‌شده نیست")
        if code in items:
            warnings.append(f"ردیف {r}: کد تکراری {code} — ردیف قبلی جایگزین شد")
        raw = sf.cell(r, 5).value
        val = clean(sv.cell(r, 5).value)
        it = {
            "code": code, "kind": kind,
            "category": clean(sv.cell(r, 3).value),
            "title": clean(sv.cell(r, 4).value),
            "unit": clean(sv.cell(r, 6).value),
            "version": clean(sv.cell(r, 8).value),
            "note": clean(sv.cell(r, 9).value),
            "src": f"{name}!E{r}",
        }
        if isinstance(raw, str) and raw.startswith("="):
            it["formula"] = raw
            it["value"] = val
            if val is None:
                warnings.append(f"{code}: مقدار محاسبه‌شده ندارد — فایل را یک بار در اکسل باز و ذخیره کنید")
            deps = {}
            for sh, rng in refs_of(raw, name):
                if ":" in rng:
                    rows, err = read_range(wv, sh, rng)
                    deps[f"{sh}!{rng}"] = rows if err is None else err
                else:
                    if sh in wv.sheetnames:
                        deps[f"{sh}!{rng}"] = clean(wv[sh][rng].value)
            if deps:
                it["refs"] = deps
        elif kind == "جدول" and isinstance(val, str) and "!" in val:
            sh, rng = val.rsplit("!", 1)
            sh = sh.strip("'")
            rows, err = read_range(wv, sh, rng.replace("$", ""))
            if err:
                warnings.append(f"{code}: {err}")
            it["table_ref"] = val
            it["rows"] = rows
        else:
            it["value"] = val
        if kind == "آزمون":
            it["inputs"] = parse_inputs(sv.cell(r, 10).value)
            if not it["inputs"]:
                warnings.append(f"{code}: آزمون بدون ورودی (ستون J)")
        items[code] = it
    return name, items, warnings


def digest(it):
    keep = {k: v for k, v in it.items() if k != "src"}
    return hashlib.sha1(json.dumps(keep, ensure_ascii=False, sort_keys=True, default=str).encode()).hexdigest()[:12]


def main():
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    if not args:
        raise SystemExit(__doc__)
    spec = "spec"
    if "--spec" in sys.argv:
        spec = sys.argv[sys.argv.index("--spec") + 1]
        args = [a for a in args if a != spec]
    force = "--force" in sys.argv
    path = args[0]
    stem = os.path.splitext(os.path.basename(path))[0]
    # «خانواده»ی فایل: نام بدون پسوند نسخه (Pumps_v8 و Pumps_v9 یک خانواده‌اند)
    family = re.sub(r"[\s_\-.]*(v|ver|rev|r|نسخه)?[\s_\-.]*\d+$", "", stem, flags=re.I) or stem
    if "--family" in sys.argv:
        family = sys.argv[sys.argv.index("--family") + 1]
        args = [a for a in args if a != family]

    sheet, items, warnings = extract(path)

    state_path = os.path.join(spec, "state.json")
    state = {}
    if os.path.exists(state_path):
        with open(state_path, encoding="utf-8") as f:
            state = json.load(f)
    # وضعیت بر اساس کد (کدها در همه‌ی فایل‌ها یکتا هستند)
    codes = state.setdefault("codes", {})

    changes = []
    for code, it in items.items():
        h = digest(it)
        old = codes.get(code)
        if old is not None and old["h"] != h and old.get("version") == it.get("version"):
            warnings.append(f"{code}: مقدار عوض شده ولی «نسخه» همان {it.get('version')} مانده — نسخه را بالا ببرید")
        if force or old is None or old["h"] != h:
            changes.append({"change": "new" if old is None else "changed",
                            **({"prev_version": old["version"]} if old else {}), **it})
    for code, old in codes.items():
        if old.get("family") == family and code not in items:
            changes.append({"change": "removed", "code": code, "prev_version": old["version"],
                            "note": "دیگر در فهرست قطعی نیست"})

    os.makedirs(os.path.join(spec, "pending"), exist_ok=True)
    today = datetime.date.today().isoformat()
    base = os.path.join(spec, "pending", f"{today}_{stem}")

    if not changes:
        print(f"{stem}: تغییری نسبت به آخرین استخراج نیست ({len(items)} مورد قطعی).")
        for w in warnings:
            print("هشدار:", w)
        return

    out = {"source": os.path.basename(path), "family": family, "sheet": sheet, "extracted": today,
           "final_count": len(items), "warnings": warnings, "changes": changes}
    # فشرده: هر تغییر در یک سطر (کم‌حجم برای خواندن، ولی خوانا در diff)
    with open(base + ".json", "w", encoding="utf-8") as f:
        head = {k: v for k, v in out.items() if k != "changes"}
        f.write(json.dumps(head, ensure_ascii=False, default=str)[:-1] + ', "changes": [\n')
        f.write(",\n".join(json.dumps(c, ensure_ascii=False, default=str, separators=(",", ":")) for c in changes))
        f.write("\n]}\n")

    label = {"new": "جدید", "changed": "تغییر", "removed": "حذف"}
    lines = [f"# {os.path.basename(path)} — {today}", "",
             f"{len(changes)} تغییر از {len(items)} مورد قطعی.", "",
             "| وضعیت | کد | نوع | رده | عنوان | مقدار | نسخه |", "|---|---|---|---|---|---|---|"]
    for c in changes:
        v = c.get("formula") or c.get("table_ref") or c.get("value", "")
        if isinstance(v, str):
            v = v.replace("|", "\\|")
        lines.append(f"| {label[c['change']]} | {c['code']} | {c.get('kind') or ''} | {c.get('category') or ''} | "
                     f"{c.get('title') or ''} | {v if v is not None else ''} | {c.get('version') or ''} |")
    if warnings:
        lines += ["", "## هشدارها", ""] + [f"- {w}" for w in warnings]
    with open(base + ".md", "w", encoding="utf-8") as f:
        f.write("\n".join(lines) + "\n")

    for code in [c for c, o in codes.items() if o.get("family") == family and c not in items]:
        del codes[code]
    for code, it in items.items():
        codes[code] = {"h": digest(it), "version": it.get("version"), "family": family}
    with open(state_path, "w", encoding="utf-8") as f:
        json.dump(state, f, ensure_ascii=False, indent=1, sort_keys=True)

    print(f"{stem}: {len(changes)} تغییر → {base}.json")
    for w in warnings:
        print("هشدار:", w)


if __name__ == "__main__":
    main()
